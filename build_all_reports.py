import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

docs_dir = r"D:\DTC245200328\Nam3\Ki1\UngDungAI\Student_Club_Management\docs"

def apply_table_styles(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    if len(table.rows) > 0:
        hdr_cells = table.rows[0].cells
        for cell in hdr_cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(255, 255, 255)
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

def build_doc(file_name, title, team_header, sections):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0, 51, 102)
    
    # Team Info
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_info = p_info.add_run(team_header)
    r_info.font.size = Pt(10.5)
    r_info.italic = True
    doc.add_paragraph()
    
    for sec_heading, sec_items in sections:
        if sec_heading:
            p_h = doc.add_paragraph()
            r_h = p_h.add_run(sec_heading)
            r_h.bold = True
            r_h.font.size = Pt(13)
            r_h.font.color.rgb = RGBColor(0, 102, 153)
            p_h.paragraph_format.space_before = Pt(8)
            p_h.paragraph_format.space_after = Pt(4)
        
        for item in sec_items:
            if isinstance(item, str):
                p_text = doc.add_paragraph(item)
                p_text.paragraph_format.space_after = Pt(4)
                p_text.paragraph_format.line_spacing = 1.15
            elif isinstance(item, tuple) and item[0] == 'table':
                headers, rows = item[1], item[2]
                table = doc.add_table(rows=1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                for idx, heading in enumerate(headers):
                    hdr_cells[idx].text = heading
                    for p in hdr_cells[idx].paragraphs:
                        for r in p.runs:
                            r.bold = True
                for row_data in rows:
                    row_cells = table.add_row().cells
                    for idx, cell_value in enumerate(row_data):
                        row_cells[idx].text = str(cell_value)
                doc.add_paragraph() # Spacing

    file_path = os.path.join(docs_dir, file_name)
    doc.save(file_path)
    print(f"Successfully generated: {file_name}")

# ==========================================
# 1. Project Plan
# ==========================================
p1_sections = [
    ("1. THÔNG TIN CHUNG DỰ ÁN", [
        "Đề tài: Hệ thống Quản lý Câu lạc bộ Sinh viên Tích hợp AI (Đề tài 28)",
        "Lớp / Nhóm thực hiện: Nhóm 15 - 02 Sinh viên",
        "Thành viên nhóm: La Văn Quyền (Trưởng nhóm - Leader/Dev/Architect), Nguyễn Đức Anh (Member - BA/UX/QA)",
        "Thời gian thực hiện: Từ 27/07/2026 đến 27/09/2026 (9 tuần)",
        "Kiến trúc Công nghệ: Flask + React-Vite + PostgreSQL + Docker + Pydantic + JWT"
    ]),
    ("2. BẢNG PHÂN CÔNG VAI TRÒ VÀ NĂNG LỰC NHÓM", [
        ('table', 
         ["STT", "Họ và tên", "Vai trò chính", "Nhiệm vụ phụ trách", "Ghi chú"],
         [
             ["1", "La Văn Quyền", "Trưởng nhóm / Lead Dev / Architect", "Kiến trúc hệ thống, CSDL PostgreSQL, Lập trình Backend Flask + Pydantic + JWT, Docker Compose, Tích hợp AI Engine Prompt", "Chịu trách nhiệm chung"],
             ["2", "Nguyễn Đức Anh", "Thành viên / BA / UX / QA", "Thu thập Q&A Yêu cầu, Thiết kế ScreenFlow UI/UX React-Vite, Kiểm thử chức năng & AI Prompt", "Phụ trách báo cáo & UX"]
         ]
        )
    ]),
    ("3. KẾ HOẠCH THỰC HIỆN CHI TIẾT 9 TUẦN (WBS)", [
        ('table',
         ["Mốc", "Tuần", "Công việc thực hiện (WBS)", "Sản phẩm đầu ra", "Người thực hiện"],
         [
             ["KT1", "Tuần 1", "Thu thập yêu cầu & khảo sát quy trình quản lý CLB", "File 02_requirements-qa.docx", "Nguyễn Đức Anh"],
             ["KT1", "Tuần 2", "Lập kế hoạch dự án & Viết Đặc tả Yêu cầu SRS", "File 01_project-plan & 03_requirements-spec", "La Văn Quyền & Nguyễn Đức Anh"],
             ["KT1", "Tuần 3", "Thiết kế Hướng đối tượng, ERD PostgreSQL & ScreenFlow UI", "File 04_object-oriented-design & 06_screenflow_db", "La Văn Quyền & Nguyễn Đức Anh"],
             ["KT2", "Tuần 4", "Khởi tạo dự án Flask + React-Vite, thiết lập Docker & PostgreSQL", "Mã nguồn Core Framework & Docker Compose", "La Văn Quyền"],
             ["KT2", "Tuần 5", "Lập trình CRUD Thành viên, Sự kiện, Phân quyền JWT & Điểm danh QR", "Tính năng Quản lý cơ bản", "La Văn Quyền"],
             ["KT2", "Tuần 6", "Lập trình Quản lý Nhiệm vụ (Kanban Board) & Thống kê", "Tính năng Quản lý nâng cao", "La Văn Quyền & Nguyễn Đức Anh"],
             ["KT3", "Tuần 7", "Viết Prompt Engineering & Tích hợp Gemini/OpenAI API + Local Rule Fallback", "Module AI Engine", "La Văn Quyền"],
             ["KT3", "Tuần 8", "Kiểm thử chức năng & Test Cases dữ liệu thiếu/AI edge cases", "File 05_functional-testing.docx", "Nguyễn Đức Anh"],
             ["Cuối kỳ", "Tuần 9", "Hoàn thiện Hướng dẫn sử dụng, Slide, Video Demo & Đóng gói Báo cáo", "File 07_user-guide & Trọn bộ Báo cáo", "La Văn Quyền & Nguyễn Đức Anh"]
         ]
        )
    ])
]

# Generate All
def run_all():
    team_hdr = "Nhóm 15 - La Văn Quyền (Leader) & Nguyễn Đức Anh (Member)"
    build_doc("01_GenAI_SoftwareDevelopment_project-plan.docx", "KẾ HOẠCH THỰC HIỆN DỰ ÁN\n(PROJECT PLAN)", team_hdr, p1_sections)

if __name__ == "__main__":
    run_all()
