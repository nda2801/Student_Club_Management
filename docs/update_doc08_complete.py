# -*- coding: utf-8 -*-
"""
Script to update Doc 08 (Final):
1. Fix blurry image 1.4: replace image1.png with sharp 2-column painpoint_to_solution.png and adjust aspect ratio in P20.
2. Remove 'AI' from Phân hệ 5: rename to 'Phân hệ 5: Báo cáo Thống kê' across text and diagrams.
3. Change all use case <<include>> relationships to <<extend>> across diagrams and text.
4. Replace all modified diagram images inside docx media.
"""

import os
import sys
import shutil
import zipfile
import docx
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

DOCX_FINAL = r"docs\08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang (final).docx"
DOCX_STD = r"docs\08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.docx"
MD_08 = r"docs\08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.md"

IMAGES_DIR = r"images"
KT1_DIR = r"CacGiaiDoanThucHien\GiaiDoan_KT1_KhaoSat_YeuCau_ThietKe_Tuan1_Tuan3"
CG_DIR = r"CacGiaiDoanThucHien"

# -------------------------------------------------------------
# STEP 1: MODIFY DOCX TEXT AND EXTENT USING PYTHON-DOCX
# -------------------------------------------------------------
print(">>> BƯỚC 1: CẬP NHẬT TEXT & KÍCH THƯỚC ẢNH TRONG DOCX FINAL...")
doc = docx.Document(DOCX_FINAL)

# Update P42
p42 = doc.paragraphs[42]
print("Old P42:", p42.text)
p42.text = " - Cấu trúc 5 Gói / Phân hệ (Packages): (1) Phân hệ Xác thực & Phân quyền; (2) Phân hệ Quản lý Thành viên & Ban; (3) Phân hệ Sự kiện & Điểm danh QR; (4) Phân hệ Nhiệm vụ & Kanban Board; (5) Phân hệ Báo cáo Thống kê."
print("New P42:", p42.text)

# Update P44
p44 = doc.paragraphs[44]
print("Old P44:", p44.text)
p44.text = " - Quan hệ mở rộng (Extend): UC04 extend 'Sinh mã QR độc bản'; UC09 extend 'Xác thực JWT'; UC10 extend 'RBAC Check Task Owner'; UC05 extend 'UC06 Gợi ý Phân công'; UC04 extend 'UC07 Sinh Bài đăng Thông báo'; UC04 extend 'UC08 Tóm tắt Kết quả Hoạt động'."
print("New P44:", p44.text)

# Update P61
p61 = doc.paragraphs[61]
print("Old P61:", p61.text)
p61.text = "4.3.5 Phân hệ 5: Báo cáo Thống kê:"
print("New P61:", p61.text)

# Update P63
p63 = doc.paragraphs[63]
print("Old P63:", p63.text)
p63.text = "Hình 4.3.5: Báo cáo Thống kê"
print("New P63:", p63.text)

# Update Table 7 (Use Cases Table)
t7 = doc.tables[7]
for r_idx in [3, 6, 7, 8]:
    cell_pkg = t7.rows[r_idx].cells[2]
    cell_pkg.text = "Báo cáo Thống kê"

# Also remove 'AI' prefix in use case names for UC06, UC07, UC08 in Table 7
t7.rows[6].cells[1].text = "Gợi ý Phân công Nhiệm vụ Thông minh"
t7.rows[7].cells[1].text = "Sinh Bài đăng Thông báo Sự kiện"
t7.rows[8].cells[1].text = "Tóm tắt Kết quả Hoạt động & Phản hồi"
print("  ✓ Đã cập nhật Bảng 7 Use Cases: Phân hệ 5 đổi thành 'Báo cáo Thống kê'!")

# Adjust extent of Image 1.4 in P20
p20 = doc.paragraphs[20]
# Check new image dimensions
im1 = Image.open(os.path.join(IMAGES_DIR, "painpoint_to_solution.png"))
w_px, h_px = im1.size
aspect = h_px / w_px
print(f"New painpoint_to_solution.png: {w_px}x{h_px}, aspect={aspect:.3f}")

# Target display width in Word: 5.8 inches
target_w_in = 5.8
target_h_in = target_w_in * aspect
cx_val = int(target_w_in * 914400)
cy_val = int(target_h_in * 914400)

for elem in p20._p.xpath('.//wp:extent'):
    elem.set('cx', str(cx_val))
    elem.set('cy', str(cy_val))
for elem in p20._p.xpath('.//a:ext'):
    elem.set('cx', str(cx_val))
    elem.set('cy', str(cy_val))

print(f"  ✓ Đã điều chỉnh tỉ lệ khung hình P20: cx={cx_val}, cy={cy_val} ({target_w_in:.2f}x{target_h_in:.2f} inches)")

# Save docx temporarily
temp_docx = "docs/_temp_08_mod.docx"
doc.save(temp_docx)
print("  ✓ Đã lưu thay đổi XML vào temp docx!")

# -------------------------------------------------------------
# STEP 2: REPLACE MEDIA IMAGES IN ZIP ARCHIVE
# -------------------------------------------------------------
print("\n>>> BƯỚC 2: THAY THẾ TOÀN BỘ CÁC ẢNH VÀO WORD/MEDIA/...")

media_replace_map = {
    'word/media/image1.png': os.path.join(IMAGES_DIR, 'painpoint_to_solution.png'),
    'word/media/image3.png': os.path.join(IMAGES_DIR, 'usecase_overall.png'),
    'word/media/image4.png': os.path.join(IMAGES_DIR, 'usecase_sub1_auth.png'),
    'word/media/image5.png': os.path.join(IMAGES_DIR, 'usecase_sub2_member.png'),
    'word/media/image6.png': os.path.join(IMAGES_DIR, 'usecase_sub3_event.png'),
    'word/media/image7.png': os.path.join(IMAGES_DIR, 'usecase_sub4_task.png'),
    'word/media/image8.png': os.path.join(IMAGES_DIR, 'usecase_sub5_ai.png')
}

out_final_zip = "docs/_final_with_new_media.docx"

with zipfile.ZipFile(temp_docx, 'r') as zin:
    with zipfile.ZipFile(out_final_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in media_replace_map:
                src_img = media_replace_map[item.filename]
                print(f"  -> Thay thế {item.filename} bằng {src_img} ({os.path.getsize(src_img)/1024:.1f} KB)")
                with open(src_img, 'rb') as f:
                    zout.writestr(item, f.read())
            else:
                zout.writestr(item, zin.read(item.filename))

# Overwrite DOCX_FINAL with the new content
shutil.copyfile(out_final_zip, DOCX_FINAL)
shutil.copyfile(out_final_zip, DOCX_STD)
shutil.copyfile(out_final_zip, os.path.join(CG_DIR, "08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.docx"))
shutil.copyfile(out_final_zip, os.path.join(KT1_DIR, "08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.docx"))
print("  ✓ Đã cập nhật thành công vào DOCX_FINAL và đồng bộ DOCX_STD, CacGiaiDoanThucHien!")

# Clean up temp
if os.path.exists(temp_docx):
    os.remove(temp_docx)
if os.path.exists(out_final_zip):
    os.remove(out_final_zip)

# -------------------------------------------------------------
# STEP 3: UPDATE MARKDOWN FILE 08
# -------------------------------------------------------------
print("\n>>> BƯỚC 3: CẬP NHẬT FILE 08 MARKDOWN (MD)...")

with open(MD_08, "r", encoding="utf-8") as f:
    md_content = f.read()

# Replace Phân hệ 5
md_content = md_content.replace(
    "- 4.3.5 Phân hệ 5: Trợ lý Trí tuệ Nhân tạo & Báo cáo Thống kê",
    "- 4.3.5 Phân hệ 5: Báo cáo Thống kê"
)
md_content = md_content.replace(
    'package "Phân hệ 5: Trợ lý AI & Báo cáo Thống kê"',
    'package "Phân hệ 5: Báo cáo Thống kê"'
)
md_content = md_content.replace(
    'rectangle "Phân hệ 5: Trợ lý AI & Báo cáo Thống kê"',
    'rectangle "Phân hệ 5: Báo cáo Thống kê"'
)
md_content = md_content.replace(
    "#### 4.3.5 Phân hệ 5: Trợ lý Trí tuệ Nhân tạo & Báo cáo Thống kê (AI Services & Analytics Subsystem)",
    "#### 4.3.5 Phân hệ 5: Báo cáo Thống kê (Reporting & Analytics Subsystem)"
)
md_content = md_content.replace(
    "| **Trợ lý AI & Báo cáo** |",
    "| **Báo cáo Thống kê** |"
)
md_content = md_content.replace(
    "Trợ lý AI & Báo cáo",
    "Báo cáo Thống kê"
)

# Replace all <<include>> with <<extend>> in Use Case sections
md_content = md_content.replace("<<include>>", "<<extend>>")
md_content = md_content.replace("<<Include>>", "<<Extend>>")
md_content = md_content.replace(": <<include>>", ": <<extend>>")
md_content = md_content.replace("|<<include>>|", "|<<extend>>|")

with open(MD_08, "w", encoding="utf-8") as f:
    f.write(md_content)

# Sync md to KT1
shutil.copyfile(MD_08, os.path.join(KT1_DIR, "08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.md"))
print("  ✓ Đã cập nhật và đồng bộ file 08 Markdown!")

print("\n>>> HOÀN TẤT TOÀN BỘ CÔNG VIỆC CẬP NHẬT FILE 08 FINAL VÀ HỆ THỐNG.")
