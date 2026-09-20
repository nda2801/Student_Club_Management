import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

docs_dir = r"D:\DTC245200328\Nam3\Ki1\UngDungAI\Student_Club_Management\docs"
images_dir = r"D:\DTC245200328\Nam3\Ki1\UngDungAI\Student_Club_Management\images"

md_path = os.path.join(docs_dir, "08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.md")
docx_path = os.path.join(docs_dir, "08_KhaoSat_PhanTich_YeuCau_ChucNang_PhiChucNang.docx")

# ==========================================
# 1. UPDATE MARKDOWN DOCUMENT WITH IMAGES
# ==========================================
print("Updating Markdown document...")
with open(md_path, "r", encoding="utf-8") as f:
    md_content = f.read()

# Map of sections / headers to image embeds in markdown
replacements = [
    (
        "#### B. Mã Biểu đồ PlantUML:\n\n```plantuml\n@startuml painpoint_to_solution",
        "#### B. Hình ảnh Biểu đồ & Mã Nguồn PlantUML:\n\n<p align=\"center\">\n  <img src=\"../images/painpoint_to_solution.png\" alt=\"Biểu đồ Ánh xạ Điểm đau sang Giải pháp\" width=\"85%\" />\n  <br/><em>Hình 1.4: Ánh xạ 06 Điểm đau (Pain Points) sang Giải pháp Chức năng & AI</em>\n</p>\n\n```plantuml\n@startuml painpoint_to_solution"
    ),
    (
        "#### B. Mã Biểu đồ PlantUML:\n\n```plantuml\n@startuml ai_functional_flow",
        "#### B. Hình ảnh Biểu đồ & Mã Nguồn PlantUML:\n\n<p align=\"center\">\n  <img src=\"../images/ai_functional_flow.png\" alt=\"Luồng Kiến trúc Xử lý Dữ liệu 3 Tính năng AI\" width=\"85%\" />\n  <br/><em>Hình 2.2: Luồng Kiến trúc Xử lý Dữ liệu 3 Tính năng AI (Smart Matchmaking, Generator, Summarizer)</em>\n</p>\n\n```plantuml\n@startuml ai_functional_flow"
    ),
    (
        "#### B. Mã Biểu đồ PlantUML ([`docs/plantuml/usecase_overall.puml`](file:///D:/DTC245200328/Nam3/Ki1/UngDungAI/Student_Club_Management/docs/plantuml/usecase_overall.puml)):\n\n```plantuml\n@startuml usecase_overall",
        "#### B. Hình ảnh Biểu đồ & Mã Nguồn PlantUML ([`docs/plantuml/usecase_overall.puml`](file:///D:/DTC245200328/Nam3/Ki1/UngDungAI/Student_Club_Management/docs/plantuml/usecase_overall.puml)):\n\n<p align=\"center\">\n  <img src=\"../images/usecase_overall.png\" alt=\"Biểu đồ Use Case Tổng thể Toàn Hệ thống\" width=\"95%\" />\n  <br/><em>Hình 4.2: Biểu đồ Ca Sử dụng Tổng thể (System Use Case Diagram - 4 Tác nhân, 12 Use Cases)</em>\n</p>\n\n```plantuml\n@startuml usecase_overall"
    ),
    (
        "#### 4.3.1 Phân hệ 1: Quản trị Tài khoản & Phân quyền Hệ thống (Authentication & RBAC Subsystem)\n\n```plantuml\n@startuml usecase_sub1_auth",
        "#### 4.3.1 Phân hệ 1: Quản trị Tài khoản & Phân quyền Hệ thống (Authentication & RBAC Subsystem)\n\n<p align=\"center\">\n  <img src=\"../images/usecase_sub1_auth.png\" alt=\"Use Case Phân hệ 1 - Auth & RBAC\" width=\"75%\" />\n  <br/><em>Hình 4.3.1: Use Case Phân hệ Xác thực & Phân quyền RBAC (UC01, UC12)</em>\n</p>\n\n```plantuml\n@startuml usecase_sub1_auth"
    ),
    (
        "#### 4.3.2 Phân hệ 2: Quản lý Hồ sơ Thành viên & Ban Chuyên môn (Member & Department Subsystem)\n\n```plantuml\n@startuml usecase_sub2_member",
        "#### 4.3.2 Phân hệ 2: Quản lý Hồ sơ Thành viên & Ban Chuyên môn (Member & Department Subsystem)\n\n<p align=\"center\">\n  <img src=\"../images/usecase_sub2_member.png\" alt=\"Use Case Phân hệ 2 - Member & Department\" width=\"75%\" />\n  <br/><em>Hình 4.3.2: Use Case Phân hệ Hồ sơ Thành viên, Skill Matrix & Ban Chuyên môn (UC02, UC11)</em>\n</p>\n\n```plantuml\n@startuml usecase_sub2_member"
    ),
    (
        "#### 4.3.3 Phân hệ 3: Quản lý Sự kiện & Điểm danh QR Độc bản (Event & QR Check-in Subsystem)\n\n```plantuml\n@startuml usecase_sub3_event",
        "#### 4.3.3 Phân hệ 3: Quản lý Sự kiện & Điểm danh QR Độc bản (Event & QR Check-in Subsystem)\n\n<p align=\"center\">\n  <img src=\"../images/usecase_sub3_event.png\" alt=\"Use Case Phân hệ 3 - Event & QR Check-in\" width=\"75%\" />\n  <br/><em>Hình 4.3.3: Use Case Phân hệ Quản lý Sự kiện & Điểm danh QR Độc bản (UC04, UC09)</em>\n</p>\n\n```plantuml\n@startuml usecase_sub3_event"
    ),
    (
        "#### 4.3.4 Phân hệ 4: Quản lý Nhiệm vụ & Bảng Kanban (Kanban Task Management Subsystem)\n\n```plantuml\n@startuml usecase_sub4_task",
        "#### 4.3.4 Phân hệ 4: Quản lý Nhiệm vụ & Bảng Kanban (Kanban Task Management Subsystem)\n\n<p align=\"center\">\n  <img src=\"../images/usecase_sub4_task.png\" alt=\"Use Case Phân hệ 4 - Kanban Task Management\" width=\"75%\" />\n  <br/><em>Hình 4.3.4: Use Case Phân hệ Quản lý Nhiệm vụ & Kanban Board (UC05, UC10)</em>\n</p>\n\n```plantuml\n@startuml usecase_sub4_task"
    ),
    (
        "#### 4.3.5 Phân hệ 5: Trợ lý Trí tuệ Nhân tạo & Báo cáo Thống kê (AI Services & Analytics Subsystem)\n\n```plantuml\n@startuml usecase_sub5_ai",
        "#### 4.3.5 Phân hệ 5: Trợ lý Trí tuệ Nhân tạo & Báo cáo Thống kê (AI Services & Analytics Subsystem)\n\n<p align=\"center\">\n  <img src=\"../images/usecase_sub5_ai.png\" alt=\"Use Case Phân hệ 5 - AI Services & Analytics\" width=\"75%\" />\n  <br/><em>Hình 4.3.5: Use Case Phân hệ Trợ lý AI & Báo cáo Thống kê (UC03, UC06, UC07, UC08)</em>\n</p>\n\n```plantuml\n@startuml usecase_sub5_ai"
    ),
    (
        "Mã nguồn Biểu đồ Lớp Phân tích (Class Diagram):\n\n```plantuml\n@startuml class_diagram",
        "Hình ảnh Biểu đồ Lớp Phân tích (Class Diagram):\n\n<p align=\"center\">\n  <img src=\"../images/class_diagram.png\" alt=\"UML Class Diagram\" width=\"95%\" />\n  <br/><em>Hình 5.1: Biểu đồ Lớp Phân tích & Thực thể CSDL (UML Class Diagram / Domain Entity Model)</em>\n</p>\n\n### 5.1.2 Sơ đồ Thực thể - Quan hệ Cơ sở dữ liệu (Database ERD - Crow's Foot Notation)\n\n<p align=\"center\">\n  <img src=\"../images/erd_database.png\" alt=\"Database ERD Diagram\" width=\"95%\" />\n  <br/><em>Hình 5.1.2: Sơ đồ Thực thể - Quan hệ CSDL PostgreSQL / SQLite (Crow's Foot Notation)</em>\n</p>\n\nMã nguồn Biểu đồ Lớp Phân tích (Class Diagram):\n\n```plantuml\n@startuml class_diagram"
    ),
    (
        "#### 5.2.1 Activity Diagram 01: Quy trình Đăng nhập, Xác thực JWT & Phân quyền RBAC\n\n```plantuml\n@startuml activity_auth_login",
        "#### 5.2.1 Activity Diagram 01: Quy trình Đăng nhập, Xác thực JWT & Phân quyền RBAC\n\n<p align=\"center\">\n  <img src=\"../images/activity_auth_login.png\" alt=\"Activity Diagram Auth\" width=\"85%\" />\n  <br/><em>Hình 5.2.1: Hoạt động Đăng nhập, Xác thực Bcrypt & Cấp Token JWT</em>\n</p>\n\n```plantuml\n@startuml activity_auth_login"
    ),
    (
        "#### 5.2.2 Activity Diagram 02: Quy trình Tổ chức Sự kiện & Điểm danh Tự động bằng Mã QR Độc bản\n\n```plantuml\n@startuml activity_qr_attendance",
        "#### 5.2.2 Activity Diagram 02: Quy trình Tổ chức Sự kiện & Điểm danh Tự động bằng Mã QR Độc bản\n\n<p align=\"center\">\n  <img src=\"../images/activity_qr_attendance.png\" alt=\"Activity Diagram QR Attendance\" width=\"85%\" />\n  <br/><em>Hình 5.2.2: Hoạt động Tổ chức Sự kiện & Quét QR Điểm danh qua Mobile</em>\n</p>\n\n```plantuml\n@startuml activity_qr_attendance"
    ),
    (
        "#### 5.2.3 Activity Diagram 03: Quy trình AI Smart Matchmaking với Dual-Engine Fallback\n\n```plantuml\n@startuml activity_ai_matching",
        "#### 5.2.3 Activity Diagram 03: Quy trình AI Smart Matchmaking với Dual-Engine Fallback\n\n<p align=\"center\">\n  <img src=\"../images/activity_ai_matching.png\" alt=\"Activity Diagram AI Matching\" width=\"85%\" />\n  <br/><em>Hình 5.2.3: Hoạt động AI Smart Matchmaking với cơ chế Dual-Engine Fallback (< 100ms switch)</em>\n</p>\n\n```plantuml\n@startuml activity_ai_matching"
    ),
    (
        "#### 5.3.1 Sequence Diagram 01: Xác thực Đăng nhập & Cấp JWT Token\n\n```plantuml\n@startuml sequence_auth_login",
        "#### 5.3.1 Sequence Diagram 01: Xác thực Đăng nhập & Cấp JWT Token\n\n<p align=\"center\">\n  <img src=\"../images/sequence_auth_login.png\" alt=\"Sequence Diagram Auth\" width=\"85%\" />\n  <br/><em>Hình 5.3.1: Tuần tự Xác thực Đăng nhập & Điều hướng vai trò</em>\n</p>\n\n```plantuml\n@startuml sequence_auth_login"
    ),
    (
        "#### 5.3.2 Sequence Diagram 02: Quét Mã QR Điểm danh Sự kiện Realtime\n\n```plantuml\n@startuml sequence_qr_checkin",
        "#### 5.3.2 Sequence Diagram 02: Quét Mã QR Điểm danh Sự kiện Realtime\n\n<p align=\"center\">\n  <img src=\"../images/sequence_qr_checkin.png\" alt=\"Sequence Diagram QR Checkin\" width=\"85%\" />\n  <br/><em>Hình 5.3.2: Tuần tự Quét QR Điểm danh Realtime & Chống gian lận</em>\n</p>\n\n```plantuml\n@startuml sequence_qr_checkin"
    ),
    (
        "#### 5.3.3 Sequence Diagram 03: AI Gợi ý Phân công Nhiệm vụ với Dual-Engine Fallback\n\n```plantuml\n@startuml sequence_ai_matchmaking",
        "#### 5.3.3 Sequence Diagram 03: AI Gợi ý Phân công Nhiệm vụ với Dual-Engine Fallback\n\n<p align=\"center\">\n  <img src=\"../images/sequence_ai_matchmaking.png\" alt=\"Sequence Diagram AI Matchmaking\" width=\"85%\" />\n  <br/><em>Hình 5.3.3: Tuần tự AI Gợi ý Phân công với Dual Fallback (< 100ms switch)</em>\n</p>\n\n```plantuml\n@startuml sequence_ai_matchmaking"
    ),
    (
        "#### 5.3.4 Sequence Diagram 04: Cập nhật Trạng thái Task trên Kanban Board với Kiểm soát RBAC\n\n```plantuml\n@startuml sequence_kanban_rbac",
        "#### 5.3.4 Sequence Diagram 04: Cập nhật Trạng thái Task trên Kanban Board với Kiểm soát RBAC\n\n<p align=\"center\">\n  <img src=\"../images/sequence_kanban_rbac.png\" alt=\"Sequence Diagram Kanban RBAC\" width=\"85%\" />\n  <br/><em>Hình 5.3.4: Tuần tự Cập nhật Task Kanban với Kiểm soát RBAC (HTTP 401/403)</em>\n</p>\n\n```plantuml\n@startuml sequence_kanban_rbac"
    ),
    (
        "#### 5.4.1 State Machine Diagram 01: Vòng đời Trạng thái Nhiệm vụ (Task Lifecycle)\n\n```plantuml\n@startuml state_task_lifecycle",
        "#### 5.4.1 State Machine Diagram 01: Vòng đời Trạng thái Nhiệm vụ (Task Lifecycle)\n\n<p align=\"center\">\n  <img src=\"../images/state_task_lifecycle.png\" alt=\"State Machine Task Lifecycle\" width=\"85%\" />\n  <br/><em>Hình 5.4.1: Vòng đời Trạng thái Nhiệm vụ (TO_DO -> IN_PROGRESS -> REVIEW -> DONE...)</em>\n</p>\n\n```plantuml\n@startuml state_task_lifecycle"
    ),
    (
        "#### 5.4.2 State Machine Diagram 02: Vòng đời Trạng thái Sự kiện (Activity Lifecycle)\n\n```plantuml\n@startuml state_activity_lifecycle",
        "#### 5.4.2 State Machine Diagram 02: Vòng đời Trạng thái Sự kiện (Activity Lifecycle)\n\n<p align=\"center\">\n  <img src=\"../images/state_activity_lifecycle.png\" alt=\"State Machine Activity Lifecycle\" width=\"85%\" />\n  <br/><em>Hình 5.4.2: Vòng đời Trạng thái Sự kiện (DRAFT -> PUBLISHED -> CHECKIN_ACTIVE -> COMPLETED...)</em>\n</p>\n\n```plantuml\n@startuml state_activity_lifecycle"
    ),
    (
        "### 5.5 Biểu đồ Thành phần & Triển khai (UML Component & Deployment Diagram)\n\n```plantuml\n@startuml deployment_docker",
        "### 5.5 Biểu đồ Thành phần & Triển khai (UML Component & Deployment Diagram)\n\n<p align=\"center\">\n  <img src=\"../images/deployment_docker.png\" alt=\"Deployment Diagram Docker\" width=\"95%\" />\n  <br/><em>Hình 5.5: Triển khai Kiến trúc Container 3-tier Docker Compose + Cloud AI</em>\n</p>\n\n```plantuml\n@startuml deployment_docker"
    )
]

for orig, repl in replacements:
    if orig in md_content:
        md_content = md_content.replace(orig, repl)
        print(f"Replaced: {orig[:40]}...")
    else:
        print(f"[Warning] Not found: {orig[:40]}...")

with open(md_path, "w", encoding="utf-8") as f:
    f.write(md_content)
print(f"Saved updated Markdown: {md_path}")

# ==========================================
# 2. UPDATE WORD DOCUMENT WITH EMBEDDED IMAGES
# ==========================================
print("\nUpdating Word document (Doc 08)...")
doc08 = docx.Document(docx_path)

# Map headings in Doc 08 to image files and captions
doc_image_mappings = [
    ("1.4 Ma trận Nhu cầu", "painpoint_to_solution.png", "Hình 1.4: Ánh xạ 06 Điểm đau (Pain Points) sang Giải pháp Chức năng & AI"),
    ("2.2 Nhóm Yêu cầu Chức năng Trí tuệ Nhân tạo", "ai_functional_flow.png", "Hình 2.2: Luồng Kiến trúc Xử lý Dữ liệu 3 Tính năng AI"),
    ("4.2 Biểu đồ Use Case Tổng thể", "usecase_overall.png", "Hình 4.2: Biểu đồ Ca Sử dụng Tổng thể (System Use Case Diagram)"),
    ("4.3.1 Phân hệ 1", "usecase_sub1_auth.png", "Hình 4.3.1: Use Case Phân hệ Xác thực & Phân quyền RBAC"),
    ("4.3.2 Phân hệ 2", "usecase_sub2_member.png", "Hình 4.3.2: Use Case Phân hệ Hồ sơ Thành viên & Ban Chuyên môn"),
    ("4.3.3 Phân hệ 3", "usecase_sub3_event.png", "Hình 4.3.3: Use Case Phân hệ Sự kiện & Điểm danh QR"),
    ("4.3.4 Phân hệ 4", "usecase_sub4_task.png", "Hình 4.3.4: Use Case Phân hệ Nhiệm vụ & Bảng Kanban"),
    ("4.3.5 Phân hệ 5", "usecase_sub5_ai.png", "Hình 4.3.5: Use Case Phân hệ Trợ lý AI & Báo cáo Thống kê"),
    ("5.1 Biểu đồ Lớp Phân tích", "class_diagram.png", "Hình 5.1: Biểu đồ Lớp Phân tích & Thực thể CSDL (UML Class Diagram)"),
    ("5.2.1 Activity Diagram 01", "activity_auth_login.png", "Hình 5.2.1: Hoạt động Đăng nhập, Xác thực Bcrypt & Cấp Token JWT"),
    ("5.2.2 Activity Diagram 02", "activity_qr_attendance.png", "Hình 5.2.2: Hoạt động Tổ chức Sự kiện & Quét QR Điểm danh qua Mobile"),
    ("5.2.3 Activity Diagram 03", "activity_ai_matching.png", "Hình 5.2.3: Hoạt động AI Smart Matchmaking với Dual-Engine Fallback"),
    ("5.3.1 Sequence Diagram 01", "sequence_auth_login.png", "Hình 5.3.1: Tuần tự Xác thực Đăng nhập & Điều hướng vai trò"),
    ("5.3.2 Sequence Diagram 02", "sequence_qr_checkin.png", "Hình 5.3.2: Tuần tự Quét QR Điểm danh Realtime & Chống gian lận"),
    ("5.3.3 Sequence Diagram 03", "sequence_ai_matchmaking.png", "Hình 5.3.3: Tuần tự AI Gợi ý Phân công với Dual Fallback"),
    ("5.3.4 Sequence Diagram 04", "sequence_kanban_rbac.png", "Hình 5.3.4: Tuần tự Cập nhật Task Kanban với Kiểm soát RBAC"),
    ("5.4.1 State Machine Diagram 01", "state_task_lifecycle.png", "Hình 5.4.1: Vòng đời Trạng thái Nhiệm vụ"),
    ("5.4.2 State Machine Diagram 02", "state_activity_lifecycle.png", "Hình 5.4.2: Vòng đời Trạng thái Sự kiện"),
    ("5.5 Biểu đồ Thành phần & Triển khai", "deployment_docker.png", "Hình 5.5: Triển khai Kiến trúc Container 3-tier Docker Compose + Cloud AI")
]

# Insert after code tables or headers
for heading_kw, img_fn, caption in doc_image_mappings:
    img_path = os.path.join(images_dir, img_fn)
    if not os.path.exists(img_path):
        continue
    for i, p in enumerate(doc08.paragraphs):
        if heading_kw in p.text:
            # Insert after the paragraph
            target_p = doc08.paragraphs[i+1] if i+1 < len(doc08.paragraphs) else doc08.add_paragraph()
            img_p = target_p.insert_paragraph_before()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_before = Pt(10)
            img_p.paragraph_format.space_after = Pt(3)
            r = img_p.add_run()
            r.add_picture(img_path, width=Inches(6.2))
            
            cap_p = target_p.insert_paragraph_before()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_after = Pt(10)
            rcap = cap_p.add_run(caption)
            rcap.font.name = "Times New Roman"
            rcap.italic = True
            rcap.font.size = Pt(10.5)
            rcap.font.color.rgb = RGBColor(70, 70, 70)
            print(f"Embedded {img_fn} into Doc 08 at '{heading_kw}'")
            break

# Also insert ERD into Doc 08 under Section 5.1
erd_img = os.path.join(images_dir, "erd_database.png")
if os.path.exists(erd_img):
    for i, p in enumerate(doc08.paragraphs):
        if "5.1 Biểu đồ Lớp Phân tích" in p.text:
            # Find next subheading 5.2
            for j in range(i+1, len(doc08.paragraphs)):
                if "5.2" in doc08.paragraphs[j].text:
                    target_p = doc08.paragraphs[j]
                    h_erd = target_p.insert_paragraph_before()
                    h_erd.paragraph_format.space_before = Pt(12)
                    h_erd.paragraph_format.space_after = Pt(4)
                    rh = h_erd.add_run("5.1.2 Sơ đồ Thực thể - Quan hệ Cơ sở dữ liệu (Database ERD - Crow's Foot Notation)")
                    rh.font.name = "Times New Roman"
                    rh.bold = True
                    rh.font.size = Pt(13)
                    rh.font.color.rgb = RGBColor(0, 51, 102)
                    
                    img_p = target_p.insert_paragraph_before()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(8)
                    img_p.paragraph_format.space_after = Pt(2)
                    r = img_p.add_run()
                    r.add_picture(erd_img, width=Inches(6.2))
                    
                    cap_p = target_p.insert_paragraph_before()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(10)
                    rcap = cap_p.add_run("Hình 5.1.2: Sơ đồ Thực thể - Quan hệ CSDL PostgreSQL / SQLite (Crow's Foot Notation)")
                    rcap.font.name = "Times New Roman"
                    rcap.italic = True
                    rcap.font.size = Pt(10.5)
                    rcap.font.color.rgb = RGBColor(70, 70, 70)
                    print("Embedded erd_database.png into Doc 08!")
                    break
            break

doc08.save(docx_path)
print(f"SUCCESSFULLY SAVED DOC 08 WORD DOCUMENT WITH EMBEDDED IMAGES: {docx_path}")
